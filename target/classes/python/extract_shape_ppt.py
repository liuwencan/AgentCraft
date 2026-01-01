import json
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def get_shape_info(shape):
    """提取单个形状的详细信息"""
    shape_info = {
        "name": shape.name or "(未命名)",
        "type": str(shape.shape_type).replace("MSO_SHAPE_TYPE.", ""),
        "type_code": shape.shape_type,
        "position": {
            "left": shape.left,
            "top": shape.top,
            "width": shape.width,
            "height": shape.height
        }
    }
    
    # 是否有文本框
    if hasattr(shape, "has_text_frame") and shape.has_text_frame:
        shape_info["has_text"] = True
        shape_info["text"] = shape.text_frame.text if shape.text_frame.text else ""
    else:
        shape_info["has_text"] = False
    
    # 如果是图片
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            shape_info["is_picture"] = True
            shape_info["image_ext"] = shape.image.ext
        except:
            shape_info["image_ext"] = "unknown"
    
    # 如果是组合形状
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        shape_info["is_group"] = True
        shape_info["group_children_count"] = len(shape.shapes)
    
    return shape_info


def iter_all_shapes(shapes):
    """递归遍历所有形状（包括组合内的）"""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_all_shapes(shape.shapes)


def analyze_ppt(ppt_path):
    """分析PPT文件"""
    prs = Presentation(ppt_path)
    
    result = {
        "file": ppt_path,
        "total_slides": len(prs.slides),
        "slides": []
    }
    
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_data = {
            "slide_index": slide_idx,
            "shape_count": len(list(iter_all_shapes(slide.shapes))),
            "shapes": []
        }
        
        for shape in iter_all_shapes(slide.shapes):
            shape_info = get_shape_info(shape)
            slide_data["shapes"].append(shape_info)
        
        result["slides"].append(slide_data)
    
    return result


# ========== 在这里硬编码你的PPT路径 ==========
PPT_PATH = r"E:\code\java-project\AgentCraft\src\main\resources\ppt\人工智能发展趋势深度解析（2026版）.pptx"  # 修改这里！
# ===========================================


if __name__ == "__main__":
    # 分析PPT
    try:
        print(f"正在分析: {PPT_PATH}")
        data = analyze_ppt(PPT_PATH)
        
        # 打印JSON
        json_str = json.dumps(data, ensure_ascii=False, indent=2)
        print(json_str)
        
        # 可选：保存到文件
        with open("ppt_analysis.json", "w", encoding="utf-8") as f:
            f.write(json_str)
        print("\n✅ 已保存到 ppt_analysis.json")
        
    except FileNotFoundError:
        print(f"❌ 文件不存在: {PPT_PATH}")
    except Exception as e:
        print(f"❌ 错误: {e}")